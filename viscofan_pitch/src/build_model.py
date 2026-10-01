from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule

OUT = "/home/user/IESEActivity/viscofan_pitch/Viscofan_HF_Model.xlsx"
wb = Workbook()

F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLACK = Font(name=F, color="000000", size=10)
GREEN = Font(name=F, color="008000", size=10)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=14, color="1F3864")
HDR = Font(name=F, bold=True, size=10, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="1F3864")
SECFILL = PatternFill("solid", fgColor="D9E1F2")
KEY = PatternFill("solid", fgColor="FFFF00")
NOTE = Font(name=F, italic=True, size=9, color="595959")
thin = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))

EUR = '€#,##0.0;(€#,##0.0);-'
EUR2 = '€#,##0.00;(€#,##0.00);-'
NUM = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.0x'

def put(ws, ref, val, font=BLACK, fmt=None, fill=None, comment=None):
    c = ws[ref]
    c.value = val
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if comment: c.comment = Comment(comment, "Team")
    return c

def header(ws, row, labels, start_col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=l)
        c.font = HDR; c.fill = HDRFILL
        c.alignment = Alignment(horizontal="center" if i else "left")

def section(ws, row, text, ncols=6):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SECFILL
    c = ws.cell(row=row, column=1, value=text); c.font = BOLD

def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v

# ---------------------------------------------------------------- Inputs
inp = wb.active; inp.title = "Inputs"
widths(inp, {"A": 46, "B": 13, "C": 13, "D": 13, "E": 70})
put(inp, "A1", "Viscofan (BME:VIS) – Inputs (€M unless stated)", TITLE)
put(inp, "A2", "Blue = hard-coded input · Black = formula · Green = link to another sheet · Yellow = key assumption", NOTE)

section(inp, 3, "Market data (valuation date 1 Oct 2026)", 5)
rows = [
    (4, "Share price (€)", 53.70, EUR2, BLUE, KEY, "Team: Capital IQ quote, 1 Oct 2026"),
    (5, "Diluted shares (M)", 46.20225, '#,##0.00', BLUE, None, "Capital IQ consensus avg. diluted shares FY26E (4 of 4 analysts); range 45.3M–47.0M"),
    (6, "Market capitalisation", "=B4*B5", EUR, BLACK, None, "Price × diluted shares"),
    (7, "Total debt (US$M, Capital IQ LTM)", 429.3, NUM, BLUE, None, "S&P Capital IQ Financials, LTM 30 Jun 2026 (Prompt 2, §1.2) – incl. leases and other financial liabilities"),
    (8, "Cash & equivalents (US$M, Capital IQ)", 66.0, NUM, BLUE, None, "S&P Capital IQ Financials, LTM 30 Jun 2026 (Prompt 2, §1.2)"),
    (9, "Cash & equivalents (€M, balance sheet)", 57.728, NUM, BLUE, None, "Viscofan H1 2026 Financial Supplement, balance sheet 30 Jun 2026 (Prompt 2, §1.1)"),
    (10, "Implied USD per EUR used by Capital IQ", "=B8/B9", '0.000', BLACK, None, "Cash in US$ ÷ cash in € – converts CIQ debt back to € on the same basis"),
    (11, "Total debt (€M)", "=B7/B10", EUR, BLACK, None, "CIQ total debt converted at implied rate"),
    (12, "Net debt incl. leases (€M)", "=B11-B9", EUR, BLACK, KEY, "Used for EV → equity bridge (conservative vs company 'net bank debt')"),
    (13, "Memo: net bank debt, company definition (€M)", 283.3, EUR, BLUE, None, "Viscofan 1H26 results presentation p.12 (Jun 2026); Dec 2025: €206.1M"),
    (14, "Enterprise value (€M)", "=B6+B12", EUR, BLACK, None, "Market cap + net debt"),
]
for r, l, v, fmt, font, fill, src in rows:
    put(inp, f"A{r}", l); put(inp, f"B{r}", v, font, fmt, fill); put(inp, f"E{r}", src, NOTE)

section(inp, 16, "Reported earnings (€M)", 5)
rows = [
    (17, "FY2025A revenue", 1252.0, "FY2025 earnings release 25 Feb 2026 (Prompt 2, §2.3)"),
    (18, "1H2025 revenue", 618.8, "1H26 results presentation p.3/p.14"),
    (19, "1H2026 revenue", 629.8, "1H26 results presentation p.3/p.14"),
    (20, "LTM Jun-26 revenue", "=B17-B18+B19", "FY25 − 1H25 + 1H26"),
    (21, "FY2025A EBITDA", 290.0, "FY2025 earnings release (Prompt 2, §2.3)"),
    (22, "1H2025 EBITDA", 145.2, "1H26 results presentation p.14"),
    (23, "1H2026 EBITDA", 134.5, "1H26 results presentation p.14"),
    (24, "LTM Jun-26 EBITDA", "=B21-B22+B23", "FY25 − 1H25 + 1H26"),
    (25, "LTM EBITDA margin", "=B24/B20", "Formula"),
    (26, "EV / LTM EBITDA (x)", "=B14/B24", "Cross-check: Capital IQ trailing TEV/EBITDA 9.94x (Prompt 1, §5.2)"),
    (27, "Net debt / LTM EBITDA (x)", "=B12/B24", "Formula"),
    (28, "Net bank debt / LTM EBITDA (x)", "=B13/B24", "Company definition; Dec-25: 206.1/290.0 = 0.71x"),
]
for r, l, v, src in rows:
    isf = isinstance(v, str)
    fmt = PCT if r == 25 else (MULT if r in (26, 27, 28) else EUR)
    put(inp, f"A{r}", l); put(inp, f"B{r}", v, BLACK if isf else BLUE, fmt); put(inp, f"E{r}", src, NOTE)

section(inp, 30, "Consensus – S&P Capital IQ Estimates (€)", 5)
header(inp, 31, ["", "FY2026E", "FY2027E", "FY2028E"])
put(inp, "A32", "EBITDA (€M)")
for col, v in zip("BCD", [300.54, 328.06, 355.91]): put(inp, f"{col}32", v, BLUE, EUR)
put(inp, "A33", "Dividend per share (€)")
for col, v in zip("BCD", [3.29, 3.39, 3.54]): put(inp, f"{col}33", v, BLUE, EUR2)
put(inp, "A34", "# analysts (EBITDA)")
for col in "BCD": put(inp, f"{col}34", 9, BLUE, '0')
put(inp, "E32", "Team: Capital IQ consensus in EUR (9 of 9 analysts)", NOTE)
put(inp, "E33", "Team: Capital IQ consensus (7–8 of 9). ~94% payout in FY26E → likely includes extraordinary component", NOTE)
put(inp, "A35", "Consensus target price (€)"); put(inp, "B35", 78.68, BLUE, EUR2); put(inp, "E35", "S&P Capital IQ Estimates; 10 analysts, 5 Buy / 2 Outperform / 3 Hold / 0 Sell (Prompt 1, §4.4)", NOTE)
put(inp, "A36", "Upside to consensus target"); put(inp, "B36", "=B35/B4-1", BLACK, PCT)

section(inp, 38, "Trade parameters (team assumptions)", 5)
rows = [
    (39, "Investment horizon (years)", 1.75, '0.00', "Team: 18–24 months → mid-point 21 months (exit ~mid-2028)"),
    (40, "Stop-loss (% below entry)", 0.15, PCT, "Team assumption – sits above bear-case value"),
    (41, "Stop-loss price (€)", "=B4*(1-B40)", EUR2, "Formula"),
    (42, "FX overlay: share of position hedged (sell USD fwd vs EUR)", 0.30, PCT, "Team: ≈ North America share of revenue (31%, 1H26 pres. p.5). Protects against USD weakness"),
]
for r, l, v, fmt, src in rows:
    isf = isinstance(v, str)
    put(inp, f"A{r}", l); put(inp, f"B{r}", v, BLACK if isf else BLUE, fmt, None if isf else KEY); put(inp, f"E{r}", src, NOTE)

# ---------------------------------------------------------------- Scenarios
sc = wb.create_sheet("Scenarios")
widths(sc, {"A": 44, "B": 15, "C": 15, "D": 15, "E": 60})
put(sc, "A1", "Bear / Base / Bull – exit valuation on FY2027E EBITDA (€)", TITLE)
put(sc, "A2", "Exit EV/EBITDA applied to FY27E EBITDA; net debt held flat (FY27 FCF ≈ dividends). No control premium (one share).", NOTE)
header(sc, 4, ["", "Bear", "Base", "Bull", "Driver / source"])
srows = [
    (5, "Probability", [0.15, 0.50, 0.35], PCT, "Team assumption (tilted to bull: Hormuz normalisation + H2 price increases)"),
    (6, "FY2027E EBITDA (€M)", [285, 315, 340], EUR, "Bear ≈ LTM (€279M) + modest growth; Base 4% below consensus €328M; Bull above consensus incl. Health/Pet"),
    (7, "Exit EV / EBITDA (x)", [8.0, 10.0, 11.5], MULT, "Bear ≈ peer median ~7.5–8x (Prompt 2 §6.4); Base = today's 10.0x (no re-rating); Bull = quality premium"),
]
for r, l, vals, fmt, src in srows:
    put(sc, f"A{r}", l)
    for col, v in zip("BCD", vals): put(sc, f"{col}{r}", v, BLUE, fmt, KEY)
    put(sc, f"E{r}", src, NOTE)
put(sc, "A8", "Implied enterprise value (€M)")
put(sc, "A9", "Less: net debt (€M)")
put(sc, "A10", "Equity value (€M)")
put(sc, "A11", "Diluted shares (M)")
put(sc, "A12", "Value per share (€)", BOLD)
put(sc, "A13", "Upside / (downside) vs €53.70")
put(sc, "A14", "EBITDA vs consensus FY27E")
for col in "BCD":
    put(sc, f"{col}8", f"={col}6*{col}7", BLACK, EUR)
    put(sc, f"{col}9", "=Inputs!$B$12", GREEN, EUR)
    put(sc, f"{col}10", f"={col}8-{col}9", BLACK, EUR)
    put(sc, f"{col}11", "=Inputs!$B$5", GREEN, '#,##0.00')
    c = put(sc, f"{col}12", f"={col}10/{col}11", BOLD, EUR2); c.border = TOP
    put(sc, f"{col}13", f"={col}12/Inputs!$B$4-1", BLACK, PCT)
    put(sc, f"{col}14", f"={col}6/Inputs!$C$32-1", BLACK, PCT)

section(sc, 16, "Trade summary", 5)
summ = [
    (17, "Probability-weighted target price (€)", "=SUMPRODUCT(B5:D5,B12:D12)", EUR2, "Headline 18–24m target price"),
    (18, "Price upside to target", "=B17/Inputs!B4-1", PCT, ""),
    (19, "Probability check (must be 100%)", "=SUM(B5:D5)", PCT, ""),
    (20, "Upside / downside ratio (bull vs bear)", "=(D12-Inputs!B4)/(Inputs!B4-B12)", MULT, "Upside to bull ÷ downside to bear"),
    (21, "Dividends received over horizon (€/sh)", "=Inputs!B33+Inputs!C33", EUR2, "FY26E + FY27E consensus DPS (interim Dec + final Jun); may include extraordinary"),
    (22, "Expected total shareholder return", "=(B17+B21)/Inputs!B4-1", PCT, ""),
    (23, "Annualised expected TSR", "=(1+B22)^(1/Inputs!B39)-1", PCT, ""),
    (24, "Stop-loss price (€)", "=Inputs!B41", EUR2, "Exit if breached without thesis-confirming news"),
    (25, "Cross-check: base case at consensus FY27E EBITDA (€)", "=(Inputs!C32*C7-Inputs!B12)/Inputs!B5", EUR2, "Base multiple × consensus EBITDA"),
    (26, "Cross-check: EV/EBITDA FY27E implied by consensus TP (x)", "=(Inputs!B35*Inputs!B5+Inputs!B12)/Inputs!C32", MULT, "What the Street's €78.68 target is paying"),
    (27, "Cross-check: DCF value per share (€)", "=DCF!B49", EUR2, "Intrinsic anchor – see DCF tab"),
]
for r, l, f, fmt, src in summ:
    put(sc, f"A{r}", l, BOLD if r in (17, 18) else BLACK)
    put(sc, f"B{r}", f, GREEN if f.startswith("=Inputs!B4") or f.startswith("=DCF") else BLACK, fmt, KEY if r == 17 else None)
    put(sc, f"E{r}", src, NOTE)

section(sc, 29, "Scenario narratives", 5)
narr = [
    ("Middle East / energy", "Cost shock persists into 2027", "Gradual normalisation; H2-26 price increases stick", "Hormuz deal; gas & plastics fall back"),
    ("EUR/USD", "USD weakens again (key risk)", "Broadly flat", "USD strength → FX tailwind"),
    ("Volumes", "Market 2% growth", "Market 2–4% + share gains (cellulose)", "Above-market (gut replacement, Zacapu)"),
    ("Health & Pet Treats", "Ignored", "Small contribution", "Re-rated as growth option"),
]
for i, (k, b, m, u) in enumerate(narr):
    r = 30 + i
    put(sc, f"A{r}", k); put(sc, f"B{r}", b, NOTE); put(sc, f"C{r}", m, NOTE); put(sc, f"D{r}", u, NOTE)
    for col in "BCD": sc[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    sc.row_dimensions[r].height = 38

# ---------------------------------------------------------------- DCF
d = wb.create_sheet("DCF")
widths(d, {"A": 40, **{c: 12 for c in "BCDEFGH"}, "I": 58})
put(d, "A1", "DCF – unlevered free cash flow, € M (valuation date 1 Oct 2026)", TITLE)
put(d, "A2", "End-year discounting from valuation date; Q4-2026 cash flow excluded (conservative). FY25A actuals for reference.", NOTE)
years = ["FY2025A", "FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E", "FY2031E"]
cols = list("BCDEFGH")
header(d, 4, ["Assumptions"] + years)
put(d, "A5", "Revenue growth")
for c, v in zip(cols[1:], [0.03, 0.05, 0.05, 0.045, 0.04, 0.035]): put(d, f"{c}5", v, BLUE, PCT)
put(d, "I5", "FY26: 1H +1.8% reported, H2 helped by price increases/less FX; FY27–28 market 2–4% vol + price/mix + Zacapu", NOTE)
put(d, "A6", "EBITDA margin")
put(d, "B6", "=B12/B11", BLACK, PCT)
for c, v in zip(cols[1:], [0.217, 0.2325, 0.2375, 0.24, 0.24, 0.24]): put(d, f"{c}6", v, BLUE, PCT, KEY)
put(d, "I6", "FY26 ≈ guidance low end; FY27 recovers to like-for-like level (1H26 LFL 23.1%); history 21.5–25%", NOTE)
put(d, "A7", "D&A (% of revenue)")
for c in cols[1:]: put(d, f"{c}7", 0.067, BLUE, PCT)
put(d, "I7", "1H26: EBITDA 134.5 − EBIT 91.8 = 42.7 → ~6.7% of revenue", NOTE)
put(d, "A8", "Capex (€M)")
for c, v in zip(cols[1:], [100, 110, 110, 104, 108, 112]): put(d, f"{c}8", v, BLUE, EUR)
put(d, "I8", "FY26 company guidance €100M (pres. p.11); FY27–28 Beat'30 step-up (Czech plant 2028); FY29+ ≈7% of revenue", NOTE)
put(d, "A9", "Change in NWC (% of incremental revenue)"); put(d, "B9", 0.20, BLUE, PCT)
put(d, "I9", "Inventories + receivables ≈ 64% of LTM sales (Jun-26 BS) net of payables → 20% assumed", NOTE)
put(d, "A10", "Tax rate on EBIT"); put(d, "B10", 0.22, BLUE, PCT)
put(d, "I10", "1H26 effective rate 21.5% (19.5 / 90.5)", NOTE)

header(d, 11 - 0, [""] * 0)  # no-op
lines = {11: "Revenue", 12: "EBITDA", 13: "D&A", 14: "EBIT", 15: "Taxes on EBIT", 16: "NOPAT",
         17: "+ D&A", 18: "− Capex", 19: "− Increase in NWC", 20: "Unlevered FCF",
         21: "Discount period (years from 1 Oct 2026)", 22: "Discount factor", 23: "PV of FCF"}
for r, l in lines.items(): put(d, f"A{r}", l, BOLD if r in (11, 12, 20) else BLACK)
put(d, "B11", "=Inputs!B17", GREEN, EUR); put(d, "B12", "=Inputs!B21", GREEN, EUR)
for i, c in enumerate(cols[1:], start=1):
    p = cols[i - 1]
    put(d, f"{c}11", f"={p}11*(1+{c}5)", BLACK, EUR)
    put(d, f"{c}12", f"={c}11*{c}6", BLACK, EUR)
    put(d, f"{c}13", f"={c}11*{c}7", BLACK, EUR)
    put(d, f"{c}14", f"={c}12-{c}13", BLACK, EUR)
    put(d, f"{c}15", f"={c}14*$B$10", BLACK, EUR)
    put(d, f"{c}16", f"={c}14-{c}15", BLACK, EUR)
    put(d, f"{c}17", f"={c}13", BLACK, EUR)
    put(d, f"{c}18", f"={c}8", BLACK, EUR)
    put(d, f"{c}19", f"=({c}11-{p}11)*$B$9", BLACK, EUR)
    cc = put(d, f"{c}20", f"={c}16+{c}17-{c}18-{c}19", BOLD, EUR); cc.border = TOP
for i, c in enumerate(cols[2:]):
    if i == 0: put(d, f"{c}21", 1.25, BLUE, '0.00')
    else: put(d, f"{c}21", f"={cols[cols.index(c)-1]}21+1", BLACK, '0.00')
    put(d, f"{c}22", f"=1/(1+$B$37)^{c}21", BLACK, '0.000')
    put(d, f"{c}23", f"={c}20*{c}22", BLACK, EUR)
put(d, "I21", "FY27 cash flow received end-2027 = 1.25 years after valuation date", NOTE)
put(d, "I20", "FY26E shown for reference only – not discounted", NOTE)

section(d, 25, "WACC", 9)
w = [
    (26, "Risk-free rate", 0.0358, PCT, BLUE, "Spain 10Y bond yield 2026E (S&P Capital IQ, Prompt 2 §7.1)"),
    (27, "Levered beta", 0.85, '0.00', BLUE, "Team assumption – low-beta non-cyclical food industrial (verify CIQ beta)"),
    (28, "Equity risk premium", 0.055, PCT, BLUE, "Team assumption – European ERP"),
    (29, "Cost of equity", "=B26+B27*B28", PCT, BLACK, "CAPM"),
    (30, "Pre-tax cost of debt", 0.038, PCT, BLUE, "Euribor/SOFR-referenced bank debt (Prompt 2 §7.1)"),
    (31, "After-tax cost of debt", "=B30*(1-B10)", PCT, BLACK, ""),
    (32, "Market cap (€M)", "=Inputs!B6", EUR, GREEN, ""),
    (33, "Net debt (€M)", "=Inputs!B12", EUR, GREEN, ""),
    (34, "Equity weight", "=B32/(B32+B33)", PCT, BLACK, "Market-value weights"),
    (35, "Debt weight", "=B33/(B32+B33)", PCT, BLACK, ""),
    (36, "Terminal growth", 0.02, PCT, BLUE, "Team: ≈ long-run inflation; below casing-market volume growth (2–4%)"),
    (37, "WACC", "=B34*B29+B35*B31", PCT, BOLD, ""),
]
for r, l, v, fmt, font, src in w:
    put(d, f"A{r}", l, BOLD if r == 37 else BLACK)
    put(d, f"B{r}", v, font, fmt, KEY if r in (27, 36) else None)
    put(d, f"I{r}", src, NOTE)

section(d, 39, "Valuation", 9)
v = [
    (40, "Sum of PV of FCF FY27–31", "=SUM(D23:H23)", EUR),
    (41, "Terminal-year FCF × (1+g)", "=H20*(1+B36)", EUR),
    (42, "Terminal value (Gordon growth)", "=B41/(B37-B36)", EUR),
    (43, "PV of terminal value", "=B42*H22", EUR),
    (44, "Enterprise value", "=B40+B43", EUR),
    (45, "Terminal value as % of EV", "=B43/B44", PCT),
    (46, "Implied exit EV/EBITDA FY31 (x)", "=B42/H12", MULT),
    (47, "Less: net debt", "=Inputs!B12", EUR),
    (48, "Equity value", "=B44-B47", EUR),
    (49, "DCF value per share (€)", "=B48/Inputs!B5", EUR2),
    (50, "Upside / (downside) vs current price", "=B49/Inputs!B4-1", PCT),
]
for r, l, f, fmt in v:
    put(d, f"A{r}", l, BOLD if r in (44, 49) else BLACK)
    put(d, f"B{r}", f, GREEN if f.startswith("=Inputs") else (BOLD if r == 49 else BLACK), fmt, KEY if r == 49 else None)

# ---------------------------------------------------------------- Sensitivity
s = wb.create_sheet("Sensitivity")
widths(s, {"A": 22, **{c: 11 for c in "BCDEFGHIJ"}})
put(s, "A1", "Sensitivities – value per share (€); green = above €53.70, red = below", TITLE)
put(s, "A3", "1) DCF: WACC (rows) × terminal growth (columns)", BOLD)
put(s, "A4", "WACC \\ g", BOLD)
gs = [0.010, 0.015, 0.020, 0.025, 0.030]
ws_ = [0.065, 0.070, 0.075, 0.080, 0.085, 0.090]
for j, g in enumerate(gs): put(s, f"{'BCDEF'[j]}4", g, BLUE, PCT)
for i, wv in enumerate(ws_):
    r = 5 + i
    put(s, f"A{r}", wv, BLUE, PCT)
    for j in range(len(gs)):
        c = "BCDEF"[j]
        f = (f"=(SUMPRODUCT(DCF!$D$20:$H$20,1/(1+$A{r})^DCF!$D$21:$H$21)"
             f"+DCF!$H$20*(1+{c}$4)/($A{r}-{c}$4)/(1+$A{r})^DCF!$H$21-Inputs!$B$12)/Inputs!$B$5")
        put(s, f"{c}{r}", f, BLACK, EUR2)
rng1 = "B5:F10"
put(s, "A12", "2) Exit multiple: FY27E EBITDA €M (rows) × EV/EBITDA (columns)", BOLD)
put(s, "A13", "EBITDA \\ x", BOLD)
mults = [8.0, 9.0, 10.0, 11.0, 12.0]
ebs = [280, 290, 300, 315, 328.06, 340, 355]
for j, m in enumerate(mults): put(s, f"{'BCDEF'[j]}13", m, BLUE, MULT)
for i, e in enumerate(ebs):
    r = 14 + i
    put(s, f"A{r}", e, BLUE, EUR)
    for j in range(len(mults)):
        c = "BCDEF"[j]
        put(s, f"{c}{r}", f"=($A{r}*{c}$13-Inputs!$B$12)/Inputs!$B$5", BLACK, EUR2)
put(s, "A22", "€328.06M = consensus FY27E; €315M = team base", NOTE)
rng2 = "B14:F20"
gfill = PatternFill("solid", fgColor="C6EFCE"); rfill = PatternFill("solid", fgColor="FFC7CE")
for rng in (rng1, rng2):
    s.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["Inputs!$B$4"], fill=gfill))
    s.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["Inputs!$B$4"], fill=rfill))

# ---------------------------------------------------------------- 1H26 data
h = wb.create_sheet("1H26 Data")
widths(h, {"A": 38, "B": 12, "C": 12, "D": 12, "E": 12, "F": 50})
put(h, "A1", "1H26 results – key data (€M) · Source: Viscofan Results Presentation Jan–Jun 2026", TITLE)
header(h, 3, ["P&L", "1H26", "1H25", "% chg", "LFL % chg", "Source"])
pl = [("Revenue", 629.8, 618.8, 0.048, "p.14"), ("EBITDA", 134.5, 145.2, 0.031, "p.14"),
      ("Operating profit (EBIT)", 91.8, 101.8, None, "p.14"), ("Net finance result", -1.3, -18.1, None, "p.10"),
      ("Profit before tax", 90.5, 83.7, None, "p.14"), ("Taxes", -19.5, -14.0, None, "p.14"),
      ("Net profit", 71.3, 69.8, None, "p.14")]
for i, (l, a, b, lfl, src) in enumerate(pl):
    r = 4 + i
    put(h, f"A{r}", l); put(h, f"B{r}", a, BLUE, NUM); put(h, f"C{r}", b, BLUE, NUM)
    put(h, f"D{r}", f"=B{r}/C{r}-1", BLACK, PCT)
    if lfl is not None: put(h, f"E{r}", lfl, BLUE, PCT)
    put(h, f"F{r}", src, NOTE)
put(h, "A11", "EBITDA margin"); put(h, "B11", "=B5/B4", BLACK, PCT); put(h, "C11", "=C5/C4", BLACK, PCT)
put(h, "D11", "=B11-C11", BLACK, '+0.0%;-0.0%'); put(h, "E11", 0.231, BLUE, PCT); put(h, "F11", "LFL margin 23.1% (p.9)", NOTE)

header(h, 13, ["Q2 (2Q26 vs 2Q25)", "2Q26", "2Q25", "% chg", "LFL % chg", "Source"])
q2 = [("Revenue", 325.0, 311.5, 0.049), ("EBITDA", 71.1, 76.3, -0.024), ("Net profit", 37.7, 38.4, None)]
for i, (l, a, b, lfl) in enumerate(q2):
    r = 14 + i
    put(h, f"A{r}", l); put(h, f"B{r}", a, BLUE, NUM); put(h, f"C{r}", b, BLUE, NUM)
    put(h, f"D{r}", f"=B{r}/C{r}-1", BLACK, PCT)
    if lfl is not None: put(h, f"E{r}", lfl, BLUE, PCT)
    put(h, f"F{r}", "p.15", NOTE)
put(h, "A17", "Consumption costs YoY: 1Q26 / 2Q26"); put(h, "B17", -0.022, BLUE, PCT); put(h, "C17", 0.168, BLUE, PCT); put(h, "F17", "p.8 – Middle East cost shock in Q2", NOTE)

header(h, 19, ["Bridges 1H25 → 1H26", "Revenue", "EBITDA", "", "", "Source"])
br = [("1H25", 618.8, 145.2), ("Like-for-like (ex energy for revenue)", 34.2, 4.5), ("Like-for-like energy", -4.6, 0),
      ("FX", -20.5, -14.6), ("Scope (Pet Mania)", 1.9, -0.6)]
for i, (l, a, b) in enumerate(br):
    r = 20 + i
    put(h, f"A{r}", l); put(h, f"B{r}", a, BLUE, NUM); put(h, f"C{r}", b, BLUE, NUM)
put(h, "A25", "1H26", BOLD); put(h, "B25", "=SUM(B20:B24)", BOLD, NUM); put(h, "C25", "=SUM(C20:C24)", BOLD, NUM)
put(h, "F20", "p.4 (revenue), p.9 (EBITDA; LFL incl. energy)", NOTE)
put(h, "A26", "FX hit as % of 1H25 base"); put(h, "B26", "=B23/B20", BLACK, PCT); put(h, "C26", "=C23/C20", BLACK, PCT)
put(h, "F26", "FX hits EBITDA ~3x harder than revenue: USD sales vs MXN/BRL/EUR costs", NOTE)

header(h, 28, ["Revenue split 1H26", "€M", "% total", "Reported", "LFL", "Source"])
seg = [("EMEA", 254.8, -0.012, -0.002), ("North America", 194.3, -0.002, 0.058), ("South America", 99.2, 0.111, 0.147),
       ("Asia-Pacific", 81.4, 0.060, 0.073), ("Food, Packaging & Ingredients", 591.7, 0.018, None),
       ("Health", 3.8, 0.573, None), ("Pet treats", 8.3, 0.878, None), ("Energy (cogeneration)", 26.0, -0.154, None)]
for i, (l, a, rep, lfl) in enumerate(seg):
    r = 29 + i
    put(h, f"A{r}", l); put(h, f"B{r}", a, BLUE, NUM); put(h, f"C{r}", f"=B{r}/Inputs!$B$19", GREEN, PCT)
    put(h, f"D{r}", rep, BLUE, PCT)
    if lfl is not None: put(h, f"E{r}", lfl, BLUE, PCT)
    put(h, f"F{r}", "p.5 (regions)" if i < 4 else "p.6 (divisions)", NOTE)
put(h, "A37", "Health + Pet as % of revenue"); put(h, "C37", "=C34+C35", BLACK, PCT)

header(h, 39, ["Cash flow & balance sheet", "€M", "", "", "", "Source"])
cf = [("Net bank debt Dec-25", 206.1), ("EBITDA", -134.5), ("Tax paid", 21.9), ("Working capital change", 41.3),
      ("Capex", 34.5), ("FX and others", 7.0), ("Shareholder remuneration", 107.0)]
for i, (l, a) in enumerate(cf):
    r = 40 + i
    put(h, f"A{r}", l); put(h, f"B{r}", a, BLUE, NUM)
put(h, "A47", "Net bank debt Jun-26", BOLD); put(h, "B47", "=SUM(B40:B46)", BOLD, NUM); put(h, "F40", "p.12 – check vs reported €283.3M", NOTE)
put(h, "A48", "Sensitivity: +10% gas price → operating profit (€M)"); put(h, "B48", -1.7, BLUE, NUM); put(h, "F48", "1H26 Half-Year Financial Report (Prompt 2, §3.2)", NOTE)
put(h, "A49", "Sensitivity: +1% interest rates → operating profit (€M)"); put(h, "B49", -1.1, BLUE, NUM); put(h, "F49", "1H26 Half-Year Financial Report (Prompt 2, §1.2)", NOTE)

# ---------------------------------------------------------------- Peers
p = wb.create_sheet("Peers")
widths(p, {"A": 34, "B": 13, "C": 13, "D": 13, "E": 13, "F": 50})
put(p, "A1", "Comparables (secondary cross-check)", TITLE)
header(p, 3, ["Company", "TEV/EBITDA", "P/E trailing", "EBITDA margin", "Net margin", "Note / source"])
peers = [("Viscofan (BME:VIS)", 9.94, 15.02, 0.2145, 0.1279, "S&P Capital IQ, priced 30 Sep 2026 (Prompt 1)"),
         ("Ebro Foods (BME:EBRO)", 7.48, 12.45, 0.1325, 0.075, "S&P Capital IQ (Prompt 1) – Spanish food benchmark")]
for i, row in enumerate(peers):
    r = 4 + i
    put(p, f"A{r}", row[0])
    for c, v, fmt in zip("BCDE", row[1:5], [MULT, MULT, PCT, PCT]): put(p, f"{c}{r}", v, BLUE, fmt)
    put(p, f"F{r}", row[5], NOTE)
for i, (n, note) in enumerate([("Shenguan Holdings (HK:0829)", "Listed collagen-casing peer – fill from Capital IQ"),
                               ("Devro / SARIA take-private (2022)", "Precedent transaction (control premium – reference only)")]):
    r = 6 + i
    put(p, f"A{r}", n)
    for c in "BCDE": put(p, f"{c}{r}", None, BLUE, None, KEY)
    put(p, f"F{r}", note, NOTE)
put(p, "A9", "Viscofan premium vs Ebro (TEV/EBITDA)"); put(p, "B9", "=B4/B5-1", BLACK, PCT)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
wb.save(OUT)
print("saved", OUT)
