import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
SRC = sys.argv[1] if len(sys.argv) > 1 else 'Accounting_Homework_Template.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'Maria_Soldevila_B_solution.xlsx'
wb=load_workbook(SRC)
g=wb['Guide']; g['C3']="Maria Soldevila (company)"; g['C4']="2011"; g['C5']="€"
t=wb['T-Accounts']
def find(title):
    for row in t.iter_rows():
        for c in row:
            if c.value==title: return c.row, c.column
    raise Exception(title)
def post(title, bb=None, dr=(), cr=(), rename=None):
    r0,c0=find(title)
    debit_normal = t.cell(r0+2,c0).value=="BB"
    if bb is not None: t.cell(r0+2, c0+1 if debit_normal else c0+2).value=bb
    for i,(ref,v) in enumerate(dr): t.cell(r0+3+i,c0).value=ref; t.cell(r0+3+i,c0+1).value=v
    for i,(ref,v) in enumerate(cr): t.cell(r0+3+i,c0+3).value=ref; t.cell(r0+3+i,c0+2).value=v
    if rename: t.cell(r0,c0).value=rename
# Beginning balances = Exhibit 1 (balance sheet at Dec 31, 2010)
post("Cash & cash equivalents",101000,[(10,850000),(12,9000)],
     [("1a",80000),("1b",160000),(3,250000),(5,162000),(6,150000),(11,30000),(13,25000)])
post("Marketable securities",20000)
post("Accounts receivable",90000,[("9a",944000)],[(10,850000),(14,10000)])
post("Inventories",32000,[(8,444000)],[("9b",439000)],rename="Finished goods inventory")
post("Prepaid rent / prepaid expenses",None,[(2,270000)],[(4,242000)],rename="Raw materials inventory")
post("Other current asset (rename)",None,[(4,242000),(5,162000),(7,40000)],[(8,444000)],rename="Work in progress (WIP)")
post("Land",300000)
post("Financial investments (LT shares)",6000,[],[("16b",6000)],rename="Website (net of amortization)")
post("Buildings",100000)
post("Accum. depreciation – Buildings",5000,[],[("16a",5000)])
post("Furniture & equipment",20000,[(13,25000)],[(12,20000)])
post("Accum. depreciation – Furn. & equip.",4000,[(12,4000)],[])
post("Software / intangible assets",None,[("1a",240000)],[],rename="Machinery")
post("Accumulated amortization",None,[],[(7,40000)],rename="Accum. depreciation – Machinery")
post("Accounts payable",50000,[(3,250000)],[(2,270000)])
post("Other liability (rename)",None,[("1b",160000)],[("1a",160000)],rename="Payable for machinery")
post("Share capital",500000)
post("Retained profits",110000,[(11,30000)],[("CE",327000)])
post("Profit & loss for the period (P&L)",None,
     [("9b",439000),(6,150000),(12,7000),(14,10000),("16a",5000),("16b",6000),("CE",327000)],
     [("9a",944000)])
j=wb['Journal']
J=[
 ("1a","Machinery installed: 240,000; 1/3 paid in cash, 2/3 deferred to end of September","Machinery",240000,None,"A+",""),
 ("1a","","Cash & cash equivalents",None,80000,"A−","I"),
 ("1a","","Payable for machinery",None,160000,"L+",""),
 ("1b","End of September: deferred 2/3 of the machinery paid","Payable for machinery",160000,None,"L−",""),
 ("1b","","Cash & cash equivalents",None,160000,"A−","I"),
 (2,"Raw materials bought on credit","Raw materials inventory",270000,None,"A+",""),
 (2,"","Accounts payable",None,270000,"L+",""),
 (3,"Payments to raw-material suppliers","Accounts payable",250000,None,"L−",""),
 (3,"","Cash & cash equivalents",None,250000,"A−","O"),
 (4,"Raw materials moved into production (NOT an expense yet)","Work in progress (WIP)",242000,None,"A+",""),
 (4,"","Raw materials inventory",None,242000,"A−",""),
 (5,"Manufacturing wages 80,000 + other manufacturing costs 82,000 → product cost","Work in progress (WIP)",162000,None,"A+",""),
 (5,"","Cash & cash equivalents",None,162000,"A−","O"),
 (6,"Selling & admin. expenses (incl. 100,000 personnel) → period expense","P&L – Selling & admin. expenses",150000,None,"Exp (OE−)",""),
 (6,"","Cash & cash equivalents",None,150000,"A−","O"),
 (7,"Machinery depreciation 240,000 / 6 years → product cost (item 7)","Work in progress (WIP)",40000,None,"A+","Non-cash"),
 (7,"","Accum. depreciation – Machinery",None,40000,"Contra-A+",""),
 (8,"600 mixers finished (242,000 + 162,000 + 40,000 = 444,000 → 740/unit); no WIP left","Finished goods inventory",444000,None,"A+",""),
 (8,"","Work in progress (WIP)",None,444000,"A−",""),
 ("9a","590 mixers sold on credit","Accounts receivable",944000,None,"A+",""),
 ("9a","","P&L – Sales revenue",None,944000,"Rev (OE+)",""),
 ("9b","COGS (FIFO): 40 × 800 + 550 × 740 = 439,000","P&L – COGS expense",439000,None,"Exp (OE−)",""),
 ("9b","","Finished goods inventory",None,439000,"A−",""),
 (10,"Collections from customers","Cash & cash equivalents",850000,None,"A+","O"),
 (10,"","Accounts receivable",None,850000,"A−",""),
 (11,"Dividends approved and paid (NOT an expense)","Retained profits",30000,None,"OE−",""),
 (11,"","Cash & cash equivalents",None,30000,"A−","F"),
 (12,"Old furniture sold for 9,000; book value 20,000 − 4,000 = 16,000 → loss 7,000","Cash & cash equivalents",9000,None,"A+","I"),
 (12,"","Accum. depreciation – Furn. & equip.",4000,None,"Contra-A−",""),
 (12,"","P&L – Loss on sale of furniture & equipment",7000,None,"Exp (OE−)",""),
 (12,"","Furniture & equipment",None,20000,"A−",""),
 (13,"New furniture & equipment bought for cash","Furniture & equipment",25000,None,"A+",""),
 (13,"","Cash & cash equivalents",None,25000,"A−","I"),
 (14,"Uncollectible receivables written down (AR shown net)","P&L – Bad debt expense",10000,None,"Exp (OE−)","Non-cash"),
 (14,"","Accounts receivable",None,10000,"A−",""),
 (15,"Securities: market 21,000 > cost 20,000 → NO entry (prudence / lower of cost or market)","(no entry)",None,None,"",""),
 ("16a","Building depreciation (selling/admin use → period expense)","P&L – Depreciation expense",5000,None,"Exp (OE−)","Non-cash"),
 ("16a","","Accum. depreciation – Buildings",None,5000,"Contra-A+",""),
 ("16b","Website amortization (reduces the net website balance to 0)","P&L – Amortization expense",6000,None,"Exp (OE−)","Non-cash"),
 ("16b","","Website (net of amortization)",None,6000,"A−",""),
 ("CE","Closing entry: net profit 2011 to retained profits","Profit & loss for the period (P&L)",327000,None,"OE−","Non-cash"),
 ("CE","","Retained profits",None,327000,"OE+",""),
]
for i,row in enumerate(J):
    r=7+i
    for col,v in zip("BCDEFGH",row):
        if v not in (None,""): j[f"{col}{r}"]=v
b=wb['Balance Sheet']; b['C4']="31 Dec 2010"; b['D4']="31 Dec 2011"
for row in b.iter_rows(min_col=2,max_col=2):
    if row[0].value=="      Intangible assets, net": row[0].value="      Machinery, net"
def setlab(ws,label,val,newlabel=None):
    for row in ws.iter_rows(min_col=2,max_col=2):
        c=row[0]
        if isinstance(c.value,str) and c.value.startswith(label):
            ws[f"C{c.row}"]=val
            if newlabel: c.value=newlabel
            return
    raise Exception(label)
s=wb['Income Statement']
setlab(s,"Sales revenue",944000,"Sales revenue (590 mixers on credit)")
setlab(s,"− Cost of goods sold",-439000,"− Cost of goods sold (FIFO: 40 × 800 + 550 × 740)")
setlab(s,"− Marketing & selling",-150000,"− Selling & administrative expenses (incl. 100,000 personnel)")
setlab(s,"− Depreciation",-11000,"− Depreciation & amortization (building 5,000 + website 6,000)")
setlab(s,"− Other operating",-10000,"− Bad debt expense (uncollectible receivables)")
setlab(s,"+/− Other non-recurring",-7000,"+/− Loss on sale of old furniture & equipment (9,000 − 16,000)")
setlab(s,"Total depreciation",51000,"Total D&A charged in 2011 incl. machinery 40,000 in product cost (positive)")
setlab(s,"Gain (+)",-7000)
c=wb['Cash Flow']
setlab(c,"+ Collections",850000)
setlab(c,"− Payments to suppliers",-250000,"− Payments to raw-material suppliers")
setlab(c,"− Payments to employees",-180000,"− Payments to employees (80,000 production + 100,000 S&A)")
setlab(c,"− Payments for other operating",-132000,"− Other manufacturing (82,000) and S&A (50,000) expenses")
setlab(c,"− Purchase of land",-240000,"− Purchase of machinery (80,000 in January + 160,000 in September)")
setlab(c,"− Purchase of buildings",-25000,"− Purchase of new furniture & equipment")
setlab(c,"+ Sale of property",9000,"+ Sale of old furniture & equipment")
setlab(c,"− Dividends",-30000)
w=wb['CF Worksheet']
rows=dict(ar=7,inv=8,prepaid=9,other_ca=10,fe=13,sw=14,fininv=15,ad_bldg=17,ad_fe=18,am_sw=19,ap=21,rp=34)
vals=dict(ar=("D",-84000,"944,000 sales − 850,000 collected − 10,000 bad debts"),
          inv=("D",-5000,"Finished goods 32,000 → 37,000 (50 units × 740)"),
          prepaid=("D",-28000,"Raw materials 270,000 bought − 242,000 used"),
          other_ca=("D",0,"No work in progress at year end"),
          fe=[("D","=7000+4000"),("E","=-25000+9000"),("H","New F&E −25,000 and sale +9,000 (CFI); loss 7,000 + AD removed 4,000 (CFO)")],
          sw=("E",-240000,"Machinery purchase (80,000 + 160,000)"),
          fininv=("D",6000,"Website amortization added back"),
          ad_bldg=("D",5000,"Building depreciation added back"),
          ad_fe=("D",-4000,"AD of the furniture sold, removed"),
          am_sw=("D",40000,"Machinery depreciation (inside product cost) added back"),
          ap=("D",20000,"Raw materials bought 270,000 − paid 250,000"),
          rp=[("D",327000),("F",-30000),("H","Net profit 327,000 (CFO) − dividends 30,000 (CFF)")])
for k,v in vals.items():
    if isinstance(v,list):
        for col,x in v: w[f"{col}{rows[k]}"]=x
    else:
        col,x,txt=v; w[f"{col}{rows[k]}"]=x; w[f"H{rows[k]}"]=txt
# ---- Corrections / report ----
cs=wb.create_sheet("Corrections",1)
cs.sheet_view.showGridLines=False
for col,wd in zip("ABCDEF",[2,6,34,34,70,2]): cs.column_dimensions[col].width=wd
F=lambda **k: Font(name="Arial",**k)
def title(r,txt):
    cs.merge_cells(f"B{r}:E{r}"); cs[f"B{r}"]=txt; cs[f"B{r}"].font=F(bold=True,color="FFFFFF")
    for col in "BCDE": cs[f"{col}{r}"].fill=PatternFill("solid",start_color="595959")
def para(r,txt,h=30,bold=False):
    cs.merge_cells(f"B{r}:E{r}"); cs[f"B{r}"]=txt; cs[f"B{r}"].font=F(size=9,bold=bold)
    cs[f"B{r}"].alignment=Alignment(wrap_text=True,vertical="top"); cs.row_dimensions[r].height=h
cs.merge_cells("B1:E1"); cs['B1']="MARIA SOLDEVILA (B) – CORRECTIONS REPORT"; cs['B1'].font=F(bold=True,size=14,color="FFFFFF")
for col in "BCDE": cs[f"{col}1"].fill=PatternFill("solid",start_color="404040")
r=3; title(r,"1. About the file you sent (Maria_Soldevila_B.ods)")
r+=1; para(r,"It contained no 2011 work. It is the (A) solution file unchanged (period 2010, entries 09-1 to CE of 2010). Everything for (B) is therefore built from scratch in this workbook, starting from the professor's Exhibit 1.",32)
r+=2; title(r,"2. Opening balances: Exhibit 1 differs from the earlier (A) solution")
rows_=[("F&E accumulated depreciation","0 (assumed no depreciation)","4,000 (20,000 / 5 years)"),
       ("Website","12,000 (assumed no amortization)","6,000 net (12,000 / 2 years)"),
       ("Retained profits (2010 net profit)","120,000","110,000"),
       ("Total assets at Dec 31, 2010","670,000","660,000")]
r+=1
for i,h in enumerate(["","Item","(A) solution had","Exhibit 1 (correct)"]):
    cs[f"{'BCDE'[i]}{r}"]=h; cs[f"{'BCDE'[i]}{r}"].font=F(bold=True,size=9)
for a,b_,c_ in rows_:
    r+=1; cs[f"C{r}"]=a; cs[f"D{r}"]=b_; cs[f"E{r}"]=c_
    for col in "CDE": cs[f"{col}{r}"].font=F(size=9)
    cs[f"D{r}"].fill=PatternFill("solid",start_color="FFC7CE"); cs[f"E{r}"].fill=PatternFill("solid",start_color="C6EFCE")
r+=1; para(r,"The (A) solution file has been updated to match Exhibit 1.",16)
r+=2; title(r,"3. Manufacturing: where each cost goes (the part you were unsure about)")
r+=1; para(r,"New T-accounts are needed. Three unused template Ts were renamed: RAW MATERIALS (was 'Prepaid rent'), WORK IN PROGRESS (was 'Other current asset'), MACHINERY + ACCUM. DEPRECIATION – MACHINERY (were 'Software / Accumulated amortization'). Also PAYABLE FOR MACHINERY (was 'Other liability'). 'Inventories' became FINISHED GOODS, and the website moved to the 'Financial investments' T (shown net, as in Exhibit 1).",55)
flow=[("Buy raw materials (2)","Dr Raw materials / Cr Accounts payable","It is an asset until it is used."),
      ("Use raw materials (4)","Dr WIP / Cr Raw materials","Moves into production. It is NOT an expense."),
      ("Production wages & other manufacturing costs (5)","Dr WIP / Cr Cash (162,000)","Product cost. It becomes an expense only when the mixers are SOLD."),
      ("Machinery depreciation (7)","Dr WIP / Cr Accum. depreciation – Machinery (40,000)","The uncle's instruction: production depreciation is included in the product cost."),
      ("Mixers finished (8)","Dr Finished goods / Cr WIP (444,000)","600 units → 740 per unit. No WIP left at year end."),
      ("Mixers sold (9)","Dr AR / Cr P&L Sales (944,000) and Dr P&L COGS / Cr Finished goods (439,000)","FIFO: the 40 old units at 800 first, then 550 at 740. Ending stock is 50 × 740 = 37,000."),
      ("Selling & admin. costs (6), building/website depreciation (16)","Dr P&L directly","These are period expenses, not product costs.")]
r+=1
for i,h in enumerate(["","Event","Entry","Why"]):
    cs[f"{'BCDE'[i]}{r}"]=h; cs[f"{'BCDE'[i]}{r}"].font=F(bold=True,size=9)
for a,b_,c_ in flow:
    r+=1; cs[f"C{r}"]=a; cs[f"D{r}"]=b_; cs[f"E{r}"]=c_
    for col in "CDE":
        cs[f"{col}{r}"].font=F(size=9); cs[f"{col}{r}"].alignment=Alignment(wrap_text=True,vertical="top")
    cs.row_dimensions[r].height=36
r+=2; title(r,"4. Other items that are easy to miss")
for txt in [
 "Machinery (1): only 80,000 is paid in January. The other 160,000 is a liability (Payable for machinery) until it is paid in September. Both payments are INVESTING cash flows.",
 "Dividends (11): Dr Retained profits / Cr Cash. They are NOT a P&L expense. Cash flow: FINANCING.",
 "Sale of old furniture (12): remove BOTH the cost (20,000) and its accumulated depreciation (4,000). Book value 16,000 − price 9,000 = loss 7,000. The 9,000 cash is INVESTING.",
 "New furniture (13): an asset (Dr F&E 25,000). Bought at the end of December, so no depreciation in 2011.",
 "Bad debts (14): Dr P&L bad debt expense 10,000 / Cr Accounts receivable. It is not a cash flow. AR is shown net of expected defaults.",
 "Marketable securities (15): market 21,000 > cost 20,000. By prudence (lower of cost or market) the 1,000 unrealized gain is NOT recorded. Keep them at 20,000.",
 "Depreciation (16): building 5,000 and website 6,000 go straight to the P&L. The website is then fully amortized.",
 "Taxes: ignored, as the uncle asked.",
 "Indirect cash flow: add back ALL depreciation (51,000 incl. the 40,000 inside product cost). The inventory changes then take care of the part still sitting in stock.",
]:
    r+=1; para(r,"• "+txt,28)
r+=2; title(r,"5. Results 2011")
for txt in ["Income statement: sales 944,000 − COGS 439,000 = gross profit 505,000; − S&A 150,000 − bad debts 10,000 − D&A 11,000 − loss on sale 7,000 = NET PROFIT 327,000.",
            "Balance sheet Dec 31, 2011: total 977,000 (cash 103,000, securities 20,000, AR 174,000, raw materials 28,000, finished goods 37,000, land 300,000, building net 90,000, F&E 25,000, machinery net 200,000 | AP 70,000, share capital 500,000, retained profits 407,000).",
            "Cash flow: CFO +288,000, CFI −256,000, CFF −30,000 → change +2,000 (101,000 → 103,000)."]:
    r+=1; para(r,txt,30,bold=False)
r+=2; title(r,"6. Assumptions (the case does not say)")
for txt in ["FIFO for finished goods. With weighted average (476,000 / 640 = 743.75) COGS would be 438,812.50 and ending stock 37,187.50.",
            "No 2011 depreciation on the old furniture (item 16 lists only the building and the website). If you depreciate it 4,000 before the sale, the loss becomes 3,000. Net profit is the same.",
            "The deferred 160,000 for the machinery was paid at the end of September as agreed."]:
    r+=1; para(r,"• "+txt,28)
cs.sheet_properties.tabColor="C00000"
cs.sheet_properties.pageSetUpPr.fitToPage=True; cs.page_setup.orientation="landscape"; cs.page_setup.fitToWidth=1; cs.page_setup.fitToHeight=0
wb.save(OUT)
print("saved",OUT)
