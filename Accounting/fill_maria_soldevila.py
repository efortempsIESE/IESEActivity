from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
BLUE="0000FF"
wb=load_workbook('Accounting_Homework_Template.xlsx')
g=wb['Guide']; g['C3']="Maria Soldevila (company)"; g['C4']="2010"; g['C5']="€"
t=wb['T-Accounts']
def find(title):
    for row in t.iter_rows():
        for c in row:
            if c.value==title: return c.row, c.column
    raise Exception(title)
def post(title, bb=None, dr=(), cr=(), rename=None):
    r0,c0=find(title)
    debit_normal = t.cell(r0+2,c0).value=="BB"
    if bb is not None: t.cell(r0+2, c0+1 if debit_normal else c0+2).value=bb
    for i,(ref,v) in enumerate(dr): t.cell(r0+3+i,c0).value=ref; t.cell(r0+3+i,c0+1).value=v
    for i,(ref,v) in enumerate(cr): t.cell(r0+3+i,c0+3).value=ref; t.cell(r0+3+i,c0+2).value=v
    if rename: t.cell(r0,c0).value=rename
# Opening balances = balance sheet at Dec 31, 2009 (see Journal entries 09-1..09-6)
post("Cash & cash equivalents",246000,[(2,150000),(4,220000)],[(5,210000),(6,90000),(7,80000),(8,110000),(9,5000),(10,20000)])
post("Marketable securities",None,[(10,20000)])
post("Accounts receivable",None,[("3a",310000)],[(4,220000)])
post("Inventories",24000,[(1,248000)],[("2b",80000),("3b",160000)])
post("Land",300000)
post("Buildings",100000)
post("Accum. depreciation – Buildings",None,[],[(11,5000)])
post("Accum. depreciation – Furn. & equip.",None,[],[(12,4000)])
post("Accumulated amortization",None,[],[(13,6000)])
post("Furniture & equipment",20000)
post("Software / intangible assets",12000,rename="Website (intangible asset)")
post("Accounts payable",12000,[(5,210000)],[(1,248000)])
post("Short-term loans / current portion LTD",110000,[(8,110000)],[],rename="Bank loan")
post("Other liability (rename)",80000,[(7,80000)],[],rename="Payable for land & building (B. Roca)")
post("Share capital",500000)
post("Retained profits",None,[],[("CE",110000)])
post("Profit & loss for the period (P&L)",None,
     [("2b",80000),("3b",160000),(6,90000),(9,5000),(11,5000),(12,4000),(13,6000),("CE",110000)],
     [("2a",150000),("3a",310000)])
# ---- Journal ----
j=wb['Journal']
J=[
 ("09-1","2009: partners pay in share capital (2 × 250,000) – already in BB","Cash & cash equivalents",500000,None,"A+","F"),
 ("09-1","","Share capital",None,500000,"OE+",""),
 ("09-2","2009: website developed by José Nadal, paid in cash – already in BB","Website (intangible asset)",12000,None,"A+",""),
 ("09-2","","Cash & cash equivalents",None,12000,"A−","I"),
 ("09-3","2009: bank loan received – already in BB","Cash & cash equivalents",110000,None,"A+","F"),
 ("09-3","","Bank loan",None,110000,"L+",""),
 ("09-4","2009: land 300,000 + building 100,000 bought from B. Roca – already in BB","Land",300000,None,"A+",""),
 ("09-4","","Buildings",100000,None,"A+",""),
 ("09-4","","Cash & cash equivalents",None,320000,"A−","I"),
 ("09-4","","Payable for land & building (B. Roca)",None,80000,"L+",""),
 ("09-5","2009: furniture & equipment paid in cash – already in BB","Furniture & equipment",20000,None,"A+",""),
 ("09-5","","Cash & cash equivalents",None,20000,"A−","I"),
 ("09-6","2009: 30 mixers × 800 imported, 12,000 paid – already in BB","Inventories",24000,None,"A+",""),
 ("09-6","","Cash & cash equivalents",None,12000,"A−","O"),
 ("09-6","","Accounts payable",None,12000,"L+",""),
 (1,"310 mixers imported on credit (310 × 800)","Inventories",248000,None,"A+",""),
 (1,"","Accounts payable",None,248000,"L+",""),
 ("2a","100 mixers sold for cash × 1,500","Cash & cash equivalents",150000,None,"A+","O"),
 ("2a","","P&L – Sales revenue",None,150000,"Rev (OE+)",""),
 ("2b","Cost of the 100 mixers sold (100 × 800)","P&L – COGS expense",80000,None,"Exp (OE−)",""),
 ("2b","","Inventories",None,80000,"A−",""),
 ("3a","200 mixers sold on credit × 1,550","Accounts receivable",310000,None,"A+",""),
 ("3a","","P&L – Sales revenue",None,310000,"Rev (OE+)",""),
 ("3b","Cost of the 200 mixers sold (200 × 800)","P&L – COGS expense",160000,None,"Exp (OE−)",""),
 ("3b","","Inventories",None,160000,"A−",""),
 (4,"Collections from credit customers","Cash & cash equivalents",220000,None,"A+","O"),
 (4,"","Accounts receivable",None,220000,"A−",""),
 (5,"Payment to the mixer supplier","Accounts payable",210000,None,"L−",""),
 (5,"","Cash & cash equivalents",None,210000,"A−","O"),
 (6,"Selling & administrative expenses paid in cash","P&L – Selling & admin. expenses",90000,None,"Exp (OE−)",""),
 (6,"","Cash & cash equivalents",None,90000,"A−","O"),
 (7,"Remaining 80,000 of land & building paid to B. Roca","Payable for land & building (B. Roca)",80000,None,"L−",""),
 (7,"","Cash & cash equivalents",None,80000,"A−","I"),
 (8,"Bank loan repaid","Bank loan",110000,None,"L−",""),
 (8,"","Cash & cash equivalents",None,110000,"A−","F"),
 (9,"Interest on the bank loan, paid at year end","P&L – Interest expense",5000,None,"Exp (OE−)",""),
 (9,"","Cash & cash equivalents",None,5000,"A−","O"),
 (10,"Surplus cash invested in marketable securities","Marketable securities",20000,None,"A+",""),
 (10,"","Cash & cash equivalents",None,20000,"A−","I"),
 (11,"Year-end adjustment: building depreciation 100,000 / 20 years","P&L – Depreciation expense",5000,None,"Exp (OE−)","Non-cash"),
 (11,"","Accum. depreciation – Buildings",None,5000,"Contra-A+",""),
 (12,"Year-end adjustment: furniture & equipment depreciation 20,000 / 5 years (per Exhibit 1 of case B)","P&L – Depreciation expense",4000,None,"Exp (OE−)","Non-cash"),
 (12,"","Accum. depreciation – Furn. & equip.",None,4000,"Contra-A+",""),
 (13,"Year-end adjustment: website amortization 12,000 / 2 years (per Exhibit 1 of case B)","P&L – Amortization expense",6000,None,"Exp (OE−)","Non-cash"),
 (13,"","Accumulated amortization",None,6000,"Contra-A+",""),
 ("CE","Closing entry: net profit to retained profits","Profit & loss for the period (P&L)",110000,None,"OE−","Non-cash"),
 ("CE","","Retained profits",None,110000,"OE+",""),
]
for i,row in enumerate(J):
    r=7+i
    for col,v in zip("BCDEFGH",row):
        if v not in (None,""): j[f"{col}{r}"]=v
# ---- Balance sheet headers ----
b=wb['Balance Sheet']; b['C4']="31 Dec 2009"; b['D4']="31 Dec 2010"
# ---- Income statement ----
s=wb['Income Statement']
def setlab(ws,label,val,newlabel=None):
    for row in ws.iter_rows(min_col=2,max_col=2):
        c=row[0]
        if isinstance(c.value,str) and c.value.startswith(label):
            ws[f"C{c.row}"]=val
            if newlabel: c.value=newlabel
            return
    raise Exception(label)
setlab(s,"Sales revenue",460000,"Sales revenue (100 × 1,500 + 200 × 1,550)")
setlab(s,"− Cost of goods",-240000,"− Cost of goods sold (300 × 800)")
setlab(s,"− Marketing & selling",-90000,"− Selling & administrative expenses")
setlab(s,"− Depreciation",-15000,"− Depreciation & amortization (building 5,000 + F&E 4,000 + website 6,000)")
setlab(s,"− Financial exp",-5000)
setlab(s,"− Tax",0,"− Tax expense (no tax information in the case)")
# ---- Cash flow direct ----
c=wb['Cash Flow']
setlab(c,"+ Collections",370000,"+ Collections from customers (150,000 cash sales + 220,000)")
setlab(c,"− Payments to suppliers",-210000)
setlab(c,"− Payments for other operating",-90000,"− Payments for selling & administrative expenses")
setlab(c,"− Payments for interest",-5000)
setlab(c,"− Purchase of buildings",-80000,"− Deferred payment for land & building (B. Roca)")
setlab(c,"− Investment in shares",-20000)
setlab(c,"− Loan & mortgage",-110000,"− Bank loan repayment")
# ---- CF worksheet ----
w=wb['CF Worksheet']
rows=dict(mktsec=6,ar=7,inv=8,ad_bldg=17,ad_fe=18,am_sw=19,ap=21,stloan=27,other_l=28,rp=34)
vals=dict(mktsec=("E",-20000,"Purchase of securities"),ar=("D",-90000,"Credit sales not yet collected"),
          inv=("D",-8000,"40 mixers left (32,000) vs 30 (24,000)"),ad_bldg=("D",5000,"Depreciation added back (non-cash)"),ad_fe=("D",4000,"F&E depreciation added back"),am_sw=("D",6000,"Website amortization added back"),
          ap=("D",38000,"Purchases not yet paid"),stloan=("F",-110000,"Loan repaid"),
          other_l=("E",-80000,"Deferred price of land & building paid"),rp=("D",110000,"Net profit via closing entry"))
for k,(col,v,txt) in vals.items():
    w[f"{col}{rows[k]}"]=v; w[f"H{rows[k]}"]=txt
# ---- Corrections sheet ----
cs=wb.create_sheet("Corrections",1)
cs.sheet_view.showGridLines=False
for col,wd in zip("ABCDEFG",[2,5,30,26,26,60,2]): cs.column_dimensions[col].width=wd
F=lambda **k: Font(name="Arial",**k)
cs.merge_cells("B1:F1"); cs['B1']="CORRECTIONS TO YOUR ORIGINAL FILE (ExcelCase3.ods)"; cs['B1'].font=F(bold=True,size=14,color="FFFFFF")
for col in "BCDEF": cs[f"{col}1"].fill=PatternFill("solid",start_color="404040")
cs['B1'].alignment=Alignment(horizontal="center")
hdr=["#","Where","What you had","Correct","Why"]
rows=[
 (1,"Building (Purchase Building)","1,000,000","100,000","Typo with an extra zero. Of the 400,000 paid, 300,000 is land and 100,000 is the building. This alone causes the whole 900,000 imbalance."),
 (2,"Furniture & Equipment","20,000 + 24,000 = 44,000","20,000","The 24,000 is the 30 mixers in stock at Dec 2009 (30 × 800). They are INVENTORY (a beginning balance), not furniture."),
 (3,"Inventories","248,000 (purchases only)","BB 24,000 + 248,000 − 240,000 = 32,000","Missing the 24,000 beginning stock and the cost of the 300 mixers sold (credit side). 40 mixers left × 800 = 32,000."),
 (4,"Accounts receivable + P&L sales","231,000","310,000","Credit sales = 200 mixers × 1,550 = 310,000. Total sales = 150,000 + 310,000 = 460,000. AR ends at 90,000, not 11,000."),
 (5,"P&L – COGS","missing","240,000 (80,000 + 160,000)","Each sale is two entries: revenue AND cost of goods sold (Dr P&L, Cr Inventories), 300 × 800."),
 (6,"Depreciation & amortization","empty","15,000","Building 100,000 / 20 = 5,000; furniture 20,000 / 5 = 4,000; website 12,000 / 2 = 6,000 (lives confirmed by Exhibit 1 of case B). Land is not depreciated."),
 (7,"P&L / Retained profits","P&L balance 286,000; RP empty","Net profit 110,000 closed to RP","Once COGS and depreciation are recorded, profit is 110,000. The closing entry moves it to retained profits."),
 (8,"Structure","2009 and 2010 mixed, no BB","Dec-2009 balances as BB, 2010 entries on top","The case asks for the balance sheet at Dec 31, 2009 AND at Dec 31, 2010. You need beginning balances to show both."),
 (9,"Income statement & cash flow sheets","labels only","filled (see sheets)","Net profit 110,000. CFO 65,000, CFI −100,000, CFF −110,000, change in cash −145,000 (246,000 → 101,000)."),
]
r=3
for i,h in enumerate(hdr):
    cell=cs[f"{'BCDEF'[i]}{r}"]; cell.value=h; cell.font=F(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",start_color="595959")
for row in rows:
    r+=1
    for i,v in enumerate(row):
        cell=cs[f"{'BCDEF'[i]}{r}"]; cell.value=v; cell.font=F(size=9,bold=(i==1)); cell.alignment=Alignment(wrap_text=True,vertical="top")
        if i==2: cell.fill=PatternFill("solid",start_color="FFC7CE")
        if i==3: cell.fill=PatternFill("solid",start_color="C6EFCE")
    cs.row_dimensions[r].height=48
r+=2
cs[f"B{r}"]="What you got RIGHT"; cs[f"B{r}"].font=F(bold=True)
r+=1
cs.merge_cells(f"B{r}:F{r}")
cs[f"B{r}"]=("Cash T-account (all 16 movements, ending at 101,000), accounts payable 50,000, share capital 500,000, bank loan in and out, "
             "website 12,000, land 300,000, the 80,000 obligation to B. Roca, marketable securities 20,000, S&A 90,000 and interest 5,000. "
             "Putting the 80,000 payment to B. Roca under INVESTING is also right. It is the deferred price of an asset, not a loan repayment.")
cs[f"B{r}"].font=F(size=9); cs[f"B{r}"].alignment=Alignment(wrap_text=True,vertical="top"); cs.row_dimensions[r].height=45
r+=2
cs[f"B{r}"]="Assumptions made (the case gives no data)"; cs[f"B{r}"].font=F(bold=True)
for txt in ["No income tax: the case gives no tax rate.",
            "Furniture 5-year life and website 2-year life: not stated in case (A), taken from the professor's Exhibit 1 in case (B) (balance sheet at Dec 31, 2010).",
            "Depreciation of the building is a full year (bought Dec 30, 2009).",
            "Marketable securities are treated as an investment (CFI), not a cash equivalent, because no maturity is given.",
            "The bank loan is shown as a current liability at Dec 2009 because it was repaid during 2010."]:
    r+=1; cs.merge_cells(f"B{r}:F{r}"); cs[f"B{r}"]="• "+txt; cs[f"B{r}"].font=F(size=9); cs[f"B{r}"].alignment=Alignment(wrap_text=True)
cs.sheet_properties.tabColor="C00000"
cs.sheet_properties.pageSetUpPr.fitToPage=True; cs.page_setup.orientation="landscape"; cs.page_setup.fitToWidth=1; cs.page_setup.fitToHeight=0
wb.save('Maria_Soldevila_A_solution.xlsx')
