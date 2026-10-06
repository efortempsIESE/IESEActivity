import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
SRC = sys.argv[1] if len(sys.argv) > 1 else 'Accounting_Homework_Template.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'Dermolab_solution.xlsx'
wb = load_workbook(SRC)
g = wb['Guide']; g['C3'] = "Laboratorios Dermolab"; g['C4'] = "2018–2021"; g['C5'] = "€"
t = wb['T-Accounts']
def find(title):
    for row in t.iter_rows():
        for c in row:
            if c.value == title: return c.row, c.column
    raise Exception(title)
def post(title, bb=None, dr=(), cr=(), rename=None):
    r0, c0 = find(title)
    debit_normal = t.cell(r0 + 2, c0).value == "BB"
    if bb is not None: t.cell(r0 + 2, c0 + 1 if debit_normal else c0 + 2).value = bb
    for i, (ref, v) in enumerate(dr): t.cell(r0 + 3 + i, c0).value = ref; t.cell(r0 + 3 + i, c0 + 1).value = v
    for i, (ref, v) in enumerate(cr): t.cell(r0 + 3 + i, c0 + 3).value = ref; t.cell(r0 + 3 + i, c0 + 2).value = v
    if rename: t.cell(r0, c0).value = rename

AMORT = round(1420000 / 10 * 2 / 12)          # Nov–Dec 2021 = 23,667
IMPAIR = 1420000 - AMORT - 1000000            # 396,333
# Old premises: only the balances being sold are known (BB = cost / accumulated depreciation before the sale)
post("Cash & cash equivalents", None,
     [("2", 13000000), ("4a", 600000)],
     [("3", 125000), ("4b", 900000), ("4c", 30000), ("4d", 30000), ("5", 8000), ("6", 40000), ("7a", 196600),
      ("4f", 28500), ("4g", 30000), ("8a", 1000000), ("9", 224000), ("4i", 27000), ("4j", 30000)])
post("Land", 5000000, [("1", 750000), ("3", 125000)], [("2", 5000000)])
post("Buildings", 6000000, [("4e", 630000)], [("2", 6000000)], rename="Buildings (old premises + new building)")
post("Accum. depreciation – Buildings", 5400000, [("2", 5400000)], [("4h", 31500), ("4k", 31500)])
post("Other current asset (rename)", None, [("4b", 900000), ("4c", 30000)], [("4e", 930000)],
     rename="Construction in progress (2019)")
post("Furniture & equipment", None, [("4e", 300000), ("6", 40000), ("7a", 196600), ("7b", 2000)], [],
     rename="Machinery & production facilities")
post("Accum. depreciation – Furn. & equip.", None, [], [("4h", 30000), ("4k", 34000)],
     rename="Accum. depreciation – Machinery & facilities")
post("Software / intangible assets", None, [("8a", 1000000)], [], rename="Patent – Baldnessox molecule")
post("Financial investments (LT shares)", None, [("8b", 420000)], [], rename="Development costs – Baldnessox")
post("Accumulated amortization", None, [], [("8c", AMORT), ("11", IMPAIR)],
     rename="Accum. amortization & impairment – intangibles")
post("Long-term bank loan", None, [("4d", 30000), ("4g", 30000), ("4j", 30000)], [("4a", 600000)],
     rename="Bank loan (5%, 20 annual instalments)")
post("Other liability (rename)", None, [], [("1", 750000)], rename="Deferred income – land donation (grant)")
post("Profit & loss for the period (P&L)", None,
     [("5", 8000), ("4f", 28500), ("4h", 61500), ("9", 224000), ("4i", 27000), ("8c", AMORT), ("11", IMPAIR), ("4k", 65500)],
     [("2", 7400000), ("7b", 2000), ("8b", 420000)])

j = wb['Journal']
J = [
 ("1", "Jul 2018 – land donation only PROMISED (condition: create local jobs) → no entry yet", "(no entry in 2018)", None, None, "", ""),
 ("2", "Oct 2018 – old premises sold: land 10,000,000 + building & facilities 3,000,000", "Cash & cash equivalents", 13000000, None, "A+", "I"),
 ("2", "", "Accum. depreciation – Buildings", 5400000, None, "Contra-A−", ""),
 ("2", "", "Land", None, 5000000, "A−", ""),
 ("2", "", "Buildings (old premises + new building)", None, 6000000, "A−", ""),
 ("2", "Gain: land 10M − 5M = 5M; building 3M − (6M − 5.4M) = 2.4M", "P&L – Gain on sale of non-current assets", None, 7400000, "Rev (OE+)", ""),
 ("1", "Jan 2019 – jobs created, donation formalized: land at FAIR VALUE 750,000 (not 550,000 tax value)", "Land", 750000, None, "A+", "Non-cash"),
 ("1", "", "Deferred income – land donation (grant)", None, 750000, "L+", ""),
 ("3", "Jan 2019 – demolition 25,000 + levelling & drainage 100,000 → prepare the land for use", "Land", 125000, None, "A+", ""),
 ("3", "", "Cash & cash equivalents", None, 125000, "A−", "I"),
 ("4a", "Jan 2019 – bank loan 600,000 at 5%, 20 annual instalments", "Cash & cash equivalents", 600000, None, "A+", "F"),
 ("4a", "", "Bank loan (5%, 20 annual instalments)", None, 600000, "L+", ""),
 ("4b", "2019 – building 450,000 + management & permits 150,000 + facilities 240,000 + installation 60,000", "Construction in progress (2019)", 900000, None, "A+", ""),
 ("4b", "", "Cash & cash equivalents", None, 900000, "A−", "I"),
 ("4c", "Dec 2019 – interest 600,000 × 5% CAPITALIZED (asset under construction)", "Construction in progress (2019)", 30000, None, "A+", ""),
 ("4c", "", "Cash & cash equivalents", None, 30000, "A−", "O"),
 ("4d", "Dec 2019 – 1st loan instalment 600,000 / 20", "Bank loan (5%, 20 annual instalments)", 30000, None, "L−", ""),
 ("4d", "", "Cash & cash equivalents", None, 30000, "A−", "F"),
 ("4e", "Jan 2020 – construction finished: building 600,000 + 30,000 interest; facilities 300,000", "Buildings (old premises + new building)", 630000, None, "A+", "Non-cash"),
 ("4e", "", "Machinery & production facilities", 300000, None, "A+", ""),
 ("4e", "", "Construction in progress (2019)", None, 930000, "A−", ""),
 ("5", "2020 – relocation 7,000 + transport insurance 1,000 → EXPENSE (staff 450 h × 20 = 9,000 already in salaries)", "P&L – Relocation expenses", 8000, None, "Exp (OE−)", ""),
 ("5", "", "Cash & cash equivalents", None, 8000, "A−", "O"),
 ("6", "2020 – special maintenance raising productivity 15% → improvement, CAPITALIZE", "Machinery & production facilities", 40000, None, "A+", ""),
 ("6", "", "Cash & cash equivalents", None, 40000, "A−", "I"),
 ("7a", "2020 – new equipment 200,000 − 2% discount + 600 transport", "Machinery & production facilities", 196600, None, "A+", ""),
 ("7a", "", "Cash & cash equivalents", None, 196600, "A−", "I"),
 ("7b", "2020 – installation by own staff 100 h × 20 (cost, never the 45 price)", "Machinery & production facilities", 2000, None, "A+", "Non-cash"),
 ("7b", "", "P&L – Own work capitalized", None, 2000, "Rev (OE+)", ""),
 ("4f", "Dec 2020 – interest 570,000 × 5% → EXPENSE (building already in use)", "P&L – Interest expense", 28500, None, "Exp (OE−)", ""),
 ("4f", "", "Cash & cash equivalents", None, 28500, "A−", "O"),
 ("4g", "Dec 2020 – 2nd loan instalment", "Bank loan (5%, 20 annual instalments)", 30000, None, "L−", ""),
 ("4g", "", "Cash & cash equivalents", None, 30000, "A−", "F"),
 ("4h", "Dec 2020 – depreciation: building 630,000 / 20; facilities 300,000 / 10", "P&L – Depreciation expense", 61500, None, "Exp (OE−)", "Non-cash"),
 ("4h", "", "Accum. depreciation – Buildings", None, 31500, "Contra-A+", ""),
 ("4h", "", "Accum. depreciation – Machinery & facilities", None, 30000, "Contra-A+", ""),
 ("8a", "Jan 2021 – patented molecule bought from Swedish company", "Patent – Baldnessox molecule", 1000000, None, "A+", ""),
 ("8a", "", "Cash & cash equivalents", None, 1000000, "A−", "I"),
 ("8b", "2021 – DEVELOPMENT 10,000 h × 42 → capitalize (approved, commercialized)", "Development costs – Baldnessox", 420000, None, "A+", "Non-cash"),
 ("8b", "", "P&L – Own work capitalized", None, 420000, "Rev (OE+)", ""),
 ("9", "Jul–Dec 2021 – RESEARCH: external 224,000 → EXPENSE (17,000 h already in salaries)", "P&L – Research expense", 224000, None, "Exp (OE−)", ""),
 ("9", "", "Cash & cash equivalents", None, 224000, "A−", "O"),
 ("10", "Oct 2021 – brand value rose to 3,000,000 → internally generated brand: NO entry", "(no entry)", None, None, "", ""),
 ("4i", "Dec 2021 – interest 540,000 × 5%", "P&L – Interest expense", 27000, None, "Exp (OE−)", ""),
 ("4i", "", "Cash & cash equivalents", None, 27000, "A−", "O"),
 ("4j", "Dec 2021 – 3rd loan instalment", "Bank loan (5%, 20 annual instalments)", 30000, None, "L−", ""),
 ("4j", "", "Cash & cash equivalents", None, 30000, "A−", "F"),
 ("4k", "Dec 2021 – depreciation: building 31,500; facilities 30,000 + improvement 40,000 / 10", "P&L – Depreciation expense", 65500, None, "Exp (OE−)", "Non-cash"),
 ("4k", "", "Accum. depreciation – Buildings", None, 31500, "Contra-A+", ""),
 ("4k", "", "Accum. depreciation – Machinery & facilities", None, 34000, "Contra-A+", ""),
 ("8c", "Dec 2021 – amortization from NOVEMBER (in use): 1,420,000 / 10 × 2/12", "P&L – Amortization expense", AMORT, None, "Exp (OE−)", "Non-cash"),
 ("8c", "", "Accum. amortization & impairment – intangibles", None, AMORT, "Contra-A+", ""),
 ("11", "Dec 2021 – impairment: book 1,396,333 vs recoverable 1,000,000 (higher of 1,000,000 VIU and 800,000 offer)", "P&L – Impairment loss", IMPAIR, None, "Exp (OE−)", "Non-cash"),
 ("11", "", "Accum. amortization & impairment – intangibles", None, IMPAIR, "Contra-A+", ""),
]
for i, row in enumerate(J):
    r = 7 + i
    for col, v in zip("BCDEFGH", row):
        if v not in (None, ""): j[f"{col}{r}"] = v

# Statements are not asked for (no opening balance sheet, four years mixed) → removed
for name in ["Balance Sheet", "Income Statement", "Cash Flow", "CF Worksheet"]:
    del wb[name]

# ---- Solution sheet ----
s = wb.create_sheet("Case solution", 1)
s.sheet_view.showGridLines = False
for col, wd in zip("ABCDEFG", [2, 6, 40, 30, 62, 16, 2]): s.column_dimensions[col].width = wd
F = lambda **k: Font(name="Arial", **k)
dark, grey = PatternFill("solid", start_color="404040"), PatternFill("solid", start_color="595959")
def bar(r, txt, fill=grey, size=10):
    s.merge_cells(f"B{r}:F{r}"); s[f"B{r}"] = txt; s[f"B{r}"].font = F(bold=True, color="FFFFFF", size=size)
    for col in "BCDEF": s[f"{col}{r}"].fill = fill
def para(r, txt, h=28):
    s.merge_cells(f"B{r}:F{r}"); s[f"B{r}"] = txt; s[f"B{r}"].font = F(size=9)
    s[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top"); s.row_dimensions[r].height = h
bar(1, "DERMOLAB – NON-CURRENT ASSETS: TREATMENT OF THE 11 EVENTS", dark, 13)
r = 3
for i, h in enumerate(["#", "Event", "Treatment", "Why (IFRS)", "Amount"]):
    c = s[f"{'BCDEF'[i]}{r}"]; c.value = h; c.font = F(bold=True, color="FFFFFF"); c.fill = grey
rows = [
 ("1", "Land donated by Viladecans (promised Jul 2018, formalized Jan 2019)",
  "Jan 2019: asset LAND at fair value; credit deferred income (grant). No entry in 2018.",
  "A conditional promise isn't an asset until the condition (jobs) is met (IAS 20.7). A non-monetary grant is measured at FAIR VALUE (750,000, the expert appraisal); the 550,000 is only a tax value.", 750000),
 ("2", "Old premises sold (Oct 2018)",
  "Remove cost of land (5M) and of building (6M) and its accumulated depreciation (5.4M). Gain to P&L.",
  "Gain = price − book value. Land 10M − 5M = 5M; building 3M − 0.6M = 2.4M. Total gain 7.4M (non-recurring). Cash +13M is INVESTING.", 7400000),
 ("3", "Demolition 25,000 + levelling & drainage 100,000",
  "CAPITALIZE as part of the LAND.",
  "Costs needed to get the land ready for its intended use are part of its cost (IAS 16.16b). Land is not depreciated, so neither are these costs.", 125000),
 ("4", "New building & production facilities (2019, in use Jan 2020)",
  "During 2019: construction in progress (not depreciated). Jan 2020: transfer to Buildings 630,000 and Facilities 300,000.",
  "All costs to bring the asset into use are capitalized: building 450,000 + management & permits 150,000; facilities 240,000 + installation 60,000. 2019 loan interest (30,000) is CAPITALIZED, because a building taking a year to construct is a qualifying asset (IAS 23). From 2020 interest is an expense. Depreciation: building 630,000 / 20 = 31,500/yr; facilities 300,000 / 10 = 30,000/yr.", 930000),
 ("5", "Moving equipment to the new plant (2020)",
  "EXPENSE everything: 7,000 + 1,000 insurance = 8,000 paid. Staff 450 h: no extra entry.",
  "Relocation costs are explicitly NOT part of an asset's cost (IAS 16.20-21). The 450 h × 20 = 9,000 of staff cost is already in salaries expense; recording it again would double count. The 45/h price is never used (it includes overhead and profit).", 8000),
 ("6", "Special maintenance +15% productivity",
  "CAPITALIZE as an improvement to machinery; depreciate over 10 years (4,000/yr from 2021).",
  "It increases future economic benefits (more output), so it is not ordinary repair & maintenance, which would be an expense (IAS 16.7, 16.12).", 40000),
 ("7", "New equipment bought in 2020",
  "Cost = 200,000 − 2% (4,000) + 600 transport + own installation 100 h × 20 = 198,600.",
  "Purchase price NET of discounts plus every cost to bring it to working condition (IAS 16.16-17). Own staff at COST (20/h), recorded as 'own work capitalized' because the salaries are already expensed.", 198600),
 ("8", "Patent bought (Jan 2021) + development (10,000 h × 42)",
  "Patent 1,000,000 and development 420,000 → INTANGIBLE ASSETS. Amortize over 10 years starting NOVEMBER 2021.",
  "Acquired patent = intangible at cost (IAS 38.25). Development is capitalized when technically feasible, approved and marketable (IAS 38.57). Use the ECONOMIC life (10) when it is shorter than the legal life (15). Amortization starts when the asset is available for use (IAS 38.97): 1,420,000 / 10 × 2/12 = 23,667.", 1420000),
 ("9", "Research on two new projects (Jul–Dec 2021)",
  "EXPENSE: external 224,000. The 17,000 h (≈714,000 at 42/h) are already in salaries expense.",
  "Research costs are ALWAYS expensed (IAS 38.54). 'Too early to determine the degree of development' = still research.", 224000),
 ("10", "Brand value rose to 3,000,000",
  "NO ENTRY.",
  "Internally generated brands are never recognized as assets (IAS 38.63), and increases in value are not recorded.", 0),
 ("11", "Competitor launches a substitute (Dec 2021)",
  "IMPAIRMENT LOSS 396,333 to P&L. From 2022, amortize the new 1,000,000 over 5 years (200,000/yr).",
  "Book value 1,420,000 − 23,667 = 1,396,333. Recoverable amount = HIGHER of value in use (1,000,000) and fair value less costs to sell (800,000 offer) = 1,000,000 (IAS 36.18). Loss = 396,333. The 1,800,000 old estimate is irrelevant: an intangible is never written UP above cost.", IMPAIR),
]
for row in rows:
    r += 1
    for i, v in enumerate(row):
        c = s[f"{'BCDEF'[i]}{r}"]; c.value = v; c.font = F(size=9, bold=(i == 0))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 4: c.number_format = '#,##0;-#,##0;"-"'
    s.row_dimensions[r].height = 78
r += 2; bar(r, "P&L IMPACT BY YEAR (from these events only)")
r += 1
table = [
 ("Gain on sale of old premises (2)", 7400000, 0, 0, 0),
 ("Relocation expenses (5)", 0, 0, -8000, 0),
 ("Own work capitalized (7b, 8b)", 0, 0, 2000, 420000),
 ("Research expense – external (9)", 0, 0, 0, -224000),
 ("Interest expense (4) – 2019 capitalized", 0, 0, -28500, -27000),
 ("Depreciation building & facilities (4, 6)", 0, 0, -61500, -65500),
 ("Amortization patent & development (8)", 0, 0, 0, -AMORT),
 ("Impairment loss Baldnessox (11)", 0, 0, 0, -IMPAIR),
]
s[f"C{r}"] = "Item"; s[f"D{r}"] = "2018 / 2019"; s[f"E{r}"] = "2020"; s[f"F{r}"] = "2021"
for col in "CDEF": s[f"{col}{r}"].font = F(bold=True, size=9); s[f"{col}{r}"].fill = PatternFill("solid", start_color="F2F2F2")
first = r + 1
for lab, a, b_, c_, d in table:
    r += 1
    s[f"C{r}"] = lab
    s[f"D{r}"] = a + b_; s[f"E{r}"] = c_; s[f"F{r}"] = d
    for col in "CDEF":
        s[f"{col}{r}"].font = F(size=9)
        if col != "C": s[f"{col}{r}"].number_format = '#,##0;-#,##0;"-"'
r += 1
s[f"C{r}"] = "Total effect on profit"
for col in "DEF":
    s[f"{col}{r}"] = f"=SUM({col}{first}:{col}{r-1})"; s[f"{col}{r}"].number_format = '#,##0;-#,##0;"-"'
    s[f"{col}{r}"].border = Border(top=Side(style="thin"))
for col in "CDEF": s[f"{col}{r}"].font = F(bold=True, size=9)
r += 1; para(r, "2018 / 2019 column: everything is 2018 (the gain). 2019 has no P&L effect: the interest is capitalized and the grant is deferred. Salaries of researchers/installers are not shown; they were already recorded as personnel expense.", 30)
r += 2; bar(r, "ASSUMPTIONS & POINTS TO CONFIRM WITH THE PROFESSOR")
for txt in [
 "Land donation (1): credited to DEFERRED INCOME (IAS 20). Alternatives you may see: credit straight to income in 2019 (the only condition, the jobs, was met), or to equity 'Donations received' (Spanish GAAP). The asset side, land at 750,000, is the same in all three.",
 "Demolition (3): capitalized in the LAND. Some textbooks add it to the new building instead, which would mean depreciating it.",
 "Loan (4): instalments of 30,000 principal + interest paid each December, starting Dec 2019. All 2019 interest is capitalized in the building (the qualifying asset).",
 "Facilities (4) come into use with the building in January 2020. The improvement (6) is done at end-2020 and depreciated over 10 years from 2021.",
 "New equipment (7): no useful life is given, so no depreciation is computed.",
 "Research hours (9) are valued at the same 42/h as the development hours (8); they matter only as an explanation, because they are already in salaries.",
 "Old premises (2): their cost and accumulated depreciation are entered as beginning balances (BB) in the T-accounts, because only the sale is recorded here. The cash T-account starts at zero, so its ending balance is not the real bank balance.",
]:
    r += 1; para(r, "• " + txt, 32)
s.sheet_properties.tabColor = "C00000"
s.sheet_properties.pageSetUpPr.fitToPage = True; s.page_setup.orientation = "landscape"
s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
wb.save(OUT)
print("saved", OUT)
